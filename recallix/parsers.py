"""Document parsers. Each returns a `Parsed` with units (pages / slides / sections / line blocks).
Uploaded files are only READ and never executed."""
from __future__ import annotations

import base64
import html
import io
import re
import shutil
from dataclasses import dataclass, field

from recallix.errors import ValidationFailed, get_logger

log = get_logger("parsers")

LINES_PER_BLOCK = 80
ROWS_PER_BLOCK = 200
MIN_TEXT_CHARS_PER_PAGE = 25


@dataclass
class Unit:
    number: int           # 1-based page / slide / section / block number
    label: str            # e.g. "Page 3", "Slide 2", "Lines 81–160"
    text: str
    html: str = ""        # sanitised HTML for reconstructed views (docx / pptx)


@dataclass
class Parsed:
    file_type: str
    unit_label: str       # "Page" | "Slide" | "Section" | "Lines" | "Rows"
    units: list[Unit] = field(default_factory=list)
    needs_ocr: bool = False
    notes: list[str] = field(default_factory=list)

    @property
    def total_chars(self) -> int:
        return sum(len(u.text) for u in self.units)


# ---------------------------------------------------------------- OCR
def ocr_available() -> bool:
    return shutil.which("tesseract") is not None


def _ocr_lang() -> str:
    try:
        import pytesseract

        langs = set(pytesseract.get_languages(config=""))
    except Exception:
        return "eng"
    wanted = [l for l in ("eng", "tam") if l in langs]
    return "+".join(wanted) or "eng"


def _ocr_image(img) -> str:
    import pytesseract

    return pytesseract.image_to_string(img, lang=_ocr_lang()).strip()


# ---------------------------------------------------------------- PDF
def parse_pdf(data: bytes, ocr: bool = False, max_ocr_pages: int = 40) -> Parsed:
    import fitz  # PyMuPDF

    try:
        doc = fitz.open(stream=data, filetype="pdf")
    except Exception as exc:
        raise ValidationFailed("This PDF is corrupted or unreadable.") from exc
    if doc.needs_pass or doc.is_encrypted:
        raise ValidationFailed("This PDF is password-protected. Remove the password and upload it again.")
    if doc.page_count == 0:
        raise ValidationFailed("This PDF has no pages.")

    parsed = Parsed("pdf", "Page")
    empty_pages = 0
    for i, page in enumerate(doc):
        text = page.get_text("text").strip()
        parsed.units.append(Unit(i + 1, f"Page {i + 1}", text))
        if len(text) < MIN_TEXT_CHARS_PER_PAGE:
            empty_pages += 1

    if empty_pages > 0 and empty_pages >= doc.page_count * 0.5:
        parsed.needs_ocr = True
        if ocr:
            if not ocr_available():
                parsed.notes.append("OCR engine (Tesseract) is not installed on this server.")
            else:
                from PIL import Image

                done = 0
                for unit in parsed.units:
                    if len(unit.text) >= MIN_TEXT_CHARS_PER_PAGE:
                        continue
                    if done >= max_ocr_pages:
                        parsed.notes.append(f"OCR was limited to the first {max_ocr_pages} scanned pages.")
                        break
                    pix = doc[unit.number - 1].get_pixmap(dpi=200)
                    img = Image.open(io.BytesIO(pix.tobytes("png")))
                    unit.text = _ocr_image(img)
                    done += 1
                parsed.needs_ocr = parsed.total_chars < MIN_TEXT_CHARS_PER_PAGE * doc.page_count * 0.3
                parsed.notes.append("Text was produced by OCR and may contain recognition errors.")
    return parsed


def pdf_outline(data: bytes) -> list[tuple[int, str, int]]:
    """[(level, title, page)] from the PDF's table of contents."""
    import fitz

    try:
        with fitz.open(stream=data, filetype="pdf") as doc:
            return [(lvl, str(title)[:100], max(int(pg), 1)) for lvl, title, pg in doc.get_toc()][:300]
    except Exception:
        return []


def render_pdf_page(data: bytes, page_number: int, scale: float = 1.6) -> bytes:
    import fitz

    with fitz.open(stream=data, filetype="pdf") as doc:
        page = doc[max(0, min(page_number - 1, doc.page_count - 1))]
        return page.get_pixmap(matrix=fitz.Matrix(scale, scale)).tobytes("png")


# ---------------------------------------------------------------- DOCX
def _data_uri(blob: bytes, content_type: str) -> str:
    if len(blob) > 1_500_000 or not content_type.startswith("image/"):
        return ""
    return f"data:{content_type};base64,{base64.b64encode(blob).decode()}"


def parse_docx(data: bytes) -> Parsed:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:
        raise ValidationFailed("This Word document is corrupted or unreadable.") from exc

    parsed = Parsed("docx", "Section")
    parsed.notes.append("Word files are shown as a reconstructed layout, not the original page layout.")
    cur_text: list[str] = []
    cur_html: list[str] = []

    def flush():
        if cur_text or cur_html:
            n = len(parsed.units) + 1
            parsed.units.append(Unit(n, f"Section {n}", "\n".join(cur_text).strip(), "".join(cur_html)))
            cur_text.clear()
            cur_html.clear()

    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(child, doc)
            style = (p.style.name if p.style is not None else "") or ""
            text = p.text.strip()
            imgs = ""
            for blip in child.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}blip"):
                rid = blip.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
                try:
                    part = doc.part.related_parts[rid]
                    uri = _data_uri(part.blob, part.content_type)
                    if uri:
                        imgs += f'<img src="{uri}" style="max-width:100%">'
                except Exception:
                    continue
            is_h = style.lower().startswith(("heading", "title"))
            if is_h and text and (style.lower() in ("heading 1", "title") or sum(map(len, cur_text)) > 2500):
                flush()
            elif sum(map(len, cur_text)) > 3500 and not is_h:
                flush()
            if text:
                cur_text.append(text)
                if is_h:
                    level = re.findall(r"\d", style)
                    h = f"h{min(int(level[0]), 4) + 1}" if level else "h2"
                    cur_html.append(f"<{h}>{html.escape(text)}</{h}>")
                else:
                    cur_html.append(f"<p>{html.escape(text)}</p>")
            if imgs:
                cur_html.append(f"<p>{imgs}</p>")
        elif tag == "tbl":
            table = Table(child, doc)
            rows_html, rows_text = [], []
            for r in table.rows[:200]:
                cells = [c.text.strip() for c in r.cells]
                rows_text.append(" | ".join(cells))
                rows_html.append("<tr>" + "".join(f"<td>{html.escape(c)}</td>" for c in cells) + "</tr>")
            cur_text.append("\n".join(rows_text))
            cur_html.append(
                '<table border="1" cellpadding="6" style="border-collapse:collapse;width:100%">'
                + "".join(rows_html) + "</table>"
            )
    flush()
    if not parsed.units:
        parsed.units.append(Unit(1, "Section 1", "", ""))
    return parsed


# ---------------------------------------------------------------- PPTX
def _shape_texts(shape, out_text: list[str], out_html: list[str]) -> None:
    if shape.shape_type == 6 and hasattr(shape, "shapes"):  # group
        for s in shape.shapes:
            _shape_texts(s, out_text, out_html)
        return
    if getattr(shape, "has_table", False) and shape.has_table:
        rows = []
        for r in shape.table.rows:
            cells = [c.text.strip() for c in r.cells]
            out_text.append(" | ".join(cells))
            rows.append("<tr>" + "".join(f"<td>{html.escape(c)}</td>" for c in cells) + "</tr>")
        out_html.append('<table border="1" cellpadding="4" style="border-collapse:collapse">' + "".join(rows) + "</table>")
        return
    if shape.shape_type == 13:  # picture
        try:
            uri = _data_uri(shape.image.blob, shape.image.content_type)
            if uri:
                out_html.append(f'<p><img src="{uri}" style="max-width:100%;max-height:360px"></p>')
        except Exception:
            pass
        return
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        for para in shape.text_frame.paragraphs:
            t = "".join(r.text for r in para.runs).strip()
            if t:
                out_text.append(t)
                out_html.append(f"<p>{html.escape(t)}</p>")


def parse_pptx(data: bytes) -> Parsed:
    from pptx import Presentation

    try:
        prs = Presentation(io.BytesIO(data))
    except Exception as exc:
        raise ValidationFailed("This PowerPoint file is corrupted or unreadable.") from exc
    parsed = Parsed("pptx", "Slide")
    parsed.notes.append("Slides are shown as extracted text and images, not as the original visual design.")
    for i, slide in enumerate(prs.slides, start=1):
        texts: list[str] = []
        htmls: list[str] = []
        for shape in slide.shapes:
            try:
                _shape_texts(shape, texts, htmls)
            except Exception:
                continue
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                texts.append(f"Speaker notes: {notes}")
                htmls.append(f"<p><em>Speaker notes: {html.escape(notes)}</em></p>")
        parsed.units.append(Unit(i, f"Slide {i}", "\n".join(texts), "".join(htmls)))
    if not parsed.units:
        raise ValidationFailed("This presentation has no slides.")
    return parsed


# ---------------------------------------------------------------- text / code / rtf
def decode_text(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            return data.decode(enc)
        except (UnicodeDecodeError, UnicodeError):
            continue
    return data.decode("utf-8", errors="replace")


def parse_text(data: bytes, file_type: str) -> Parsed:
    text = decode_text(data)
    if file_type == "rtf":
        from striprtf.striprtf import rtf_to_text

        text = rtf_to_text(text)
    lines = text.splitlines()
    parsed = Parsed(file_type, "Lines")
    for start in range(0, max(len(lines), 1), LINES_PER_BLOCK):
        block = lines[start:start + LINES_PER_BLOCK]
        parsed.units.append(
            Unit(start // LINES_PER_BLOCK + 1, f"Lines {start + 1}–{start + len(block)}", "\n".join(block))
        )
    return parsed


def code_outline(text: str, ext: str) -> list[tuple[int, str]]:
    """[(line_number, 'def foo')] — functions and classes for navigation."""
    out: list[tuple[int, str]] = []
    if ext == "py":
        import ast

        try:
            tree = ast.parse(text)
            for node in ast.walk(tree):
                if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef, ast.ClassDef)):
                    kind = "class" if isinstance(node, ast.ClassDef) else "def"
                    out.append((node.lineno, f"{kind} {node.name}"))
            return sorted(out)
        except SyntaxError:
            pass
    pattern = re.compile(
        r"^\s*(?:(?:public|private|protected|static|final|async|export|default|abstract)\s+)*"
        r"(?:(class|interface|enum|struct|function|def|fn)\s+(\w+)|"
        r"[\w:<>\[\]\*&]+\s+\**(\w+)\s*\([^;]*\)\s*(?:const)?\s*\{?\s*$)"
    )
    skip = {"if", "for", "while", "switch", "return", "else", "catch"}
    for no, line in enumerate(text.splitlines(), start=1):
        m = pattern.match(line)
        if m:
            name = m.group(2) or m.group(3)
            if name and name not in skip:
                out.append((no, f"{m.group(1) or 'fn'} {name}"))
    return out[:300]


# ---------------------------------------------------------------- tables
def read_table_preview(data: bytes, file_type: str, sheet: str | None = None, max_rows: int = 50000):
    import pandas as pd

    try:
        if file_type == "csv":
            return pd.read_csv(io.BytesIO(data), nrows=max_rows, encoding_errors="replace", on_bad_lines="skip")
        return pd.read_excel(io.BytesIO(data), sheet_name=sheet or 0, nrows=max_rows, engine="openpyxl")
    except Exception as exc:
        raise ValidationFailed("This spreadsheet could not be read. It may be corrupted or encrypted.") from exc


def excel_sheets(data: bytes) -> list[str]:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        return wb.sheetnames[:30]
    finally:
        wb.close()


def parse_table(data: bytes, file_type: str, max_rows: int = 50000) -> Parsed:
    parsed = Parsed(file_type, "Rows")
    sheets = [None] if file_type == "csv" else excel_sheets(data)
    n = 0
    for sheet in sheets:
        df = read_table_preview(data, file_type, sheet, max_rows)
        name = f"{sheet}: " if sheet else ""
        header = f"Columns: {', '.join(map(str, df.columns))}\n"
        for start in range(0, max(len(df), 1), ROWS_PER_BLOCK):
            block = df.iloc[start:start + ROWS_PER_BLOCK]
            n += 1
            parsed.units.append(
                Unit(n, f"{name}Rows {start + 1}–{start + len(block)}", header + block.to_csv(index=False))
            )
    if len(parsed.units) * ROWS_PER_BLOCK >= max_rows:
        parsed.notes.append(f"Only the first {max_rows:,} rows per sheet were processed.")
    return parsed


# ---------------------------------------------------------------- images
def parse_image(data: bytes, ocr: bool = False) -> Parsed:
    from PIL import Image

    try:
        img = Image.open(io.BytesIO(data))
        img.verify()
        img = Image.open(io.BytesIO(data))
    except Exception as exc:
        raise ValidationFailed("This image is corrupted or unreadable.") from exc
    parsed = Parsed("image", "Page", [Unit(1, "Image", "")])
    parsed.needs_ocr = True
    if ocr:
        if not ocr_available():
            parsed.notes.append("OCR engine (Tesseract) is not installed on this server.")
        else:
            parsed.units[0].text = _ocr_image(img.convert("RGB"))
            parsed.needs_ocr = len(parsed.units[0].text) < MIN_TEXT_CHARS_PER_PAGE
            parsed.notes.append("Text was produced by OCR and may contain recognition errors.")
    return parsed


# ---------------------------------------------------------------- dispatcher
def parse_document(data: bytes, file_type: str, *, ocr: bool = False, max_ocr_pages: int = 40,
                   max_table_rows: int = 50000) -> Parsed:
    if file_type == "pdf":
        return parse_pdf(data, ocr, max_ocr_pages)
    if file_type == "docx":
        return parse_docx(data)
    if file_type == "pptx":
        return parse_pptx(data)
    if file_type in ("csv", "xlsx"):
        return parse_table(data, file_type, max_table_rows)
    if file_type == "image":
        return parse_image(data, ocr)
    return parse_text(data, file_type)

"""Test launcher: patches the data/AI layers with in-memory fakes, then runs the real app.py."""
import os
import runpy
from pathlib import Path

import fitz
import streamlit as st

os.environ.update({"SUPABASE_URL": "https://x.supabase.co", "SUPABASE_ANON_KEY": "a", "GEMINI_API_KEY": "k"})
from recallix import auth, db, llm, quota, storage  # noqa: E402

pdf = fitz.open()
for i in range(4):
    pdf.new_page().insert_text((72, 72), f"Chapter {i + 1}: loops repeat code while a condition is true.")
PDF = pdf.tobytes()
KIND = os.environ.get("FAKE_KIND", "pdf")
SRC = b"def add(a, b):\n    return a + b\n\nclass Cat:\n    pass\n"


def _office(kind):
    import io

    buf = io.BytesIO()
    if kind == "docx":
        from docx import Document

        d = Document()
        d.add_heading("Title", 1)
        d.add_paragraph("Hello <b>world</b>")
        d.save(buf)
    elif kind == "pptx":
        from pptx import Presentation

        prs = Presentation()
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = "Slide A"
        prs.save(buf)
    elif kind == "xlsx":
        import pandas as pd

        pd.DataFrame({"a": range(300), "b": ["x"] * 300}).to_excel(buf, index=False)
    elif kind == "csv":
        return b"a,b\n1,2\n3,4\n"
    return buf.getvalue()


BYTES = {"pdf": PDF, "code": SRC}
if KIND in ("docx", "pptx", "xlsx", "csv"):
    BYTES[KIND] = _office(KIND)
DOC = {"id": "d1", "owner_id": "u1", "original_name": {"pdf": "book.pdf", "code": "demo.py"}.get(KIND, f"file.{KIND}"),
       "file_type": KIND, "mime_type": "x", "size_bytes": 100, "storage_path": "u1/x",
       "sha256": "s", "status": "ready", "status_detail": "", "page_count": 4 if KIND == "pdf" else 1,
       "is_favourite": False, "created_at": "2026-01-01T00:00:00", "updated_at": "2026-01-01T00:00:00"}
CHUNKS = [{"chunk_index": i, "page_number": i + 1, "section": f"Page {i + 1}", "content": f"Chapter {i + 1} loops"} for i in range(4)]

if "msgs" not in st.session_state:
    st.session_state["msgs"], st.session_state["convs"] = [], []
MSGS, CONVS = st.session_state["msgs"], st.session_state["convs"]
CALLS = st.session_state.setdefault("llm_calls", [])

auth.is_logged_in = lambda: True
auth.user_id = lambda: "u1"
auth.client = lambda: None
st.session_state.setdefault("user", {"id": "u1", "email": "a@b.co", "created_at": "2026-01-01"})
db.get_prefs = lambda: {"language": "tanglish", "level": "beginner", "voice_lang": "en-IN", "auto_speak": False, "history_limit": 10}
db.list_documents = lambda limit=100, favourites_only=False: [DOC]
db.get_document = lambda i: DOC
db.list_conversations = lambda limit=100: CONVS
db.create_conversation = lambda d, t: (CONVS.append({"id": "c1", "title": t, "document_id": d, "updated_at": "2026-01-01T00:00"}) or CONVS[-1])
db.recent_messages = lambda c, n: [{"role": m["role"], "content": m["content"]} for m in MSGS][-n:]
db.add_message = lambda c, r, content, lang, sources=None: (MSGS.append({"id": len(MSGS) + 1, "role": r, "content": content, "sources": sources or []}) or MSGS[-1])
db.list_messages = lambda c, limit=200: list(MSGS)
db.touch_conversation = lambda c: None
db.record_usage = lambda k, u=1: None
db.page_chunks = lambda d, p: [c for c in CHUNKS if c["page_number"] == p]
db.all_chunks_sample = lambda d, n: CHUNKS[:n]
db.match_chunks = lambda d, e, k=6: [{**CHUNKS[2], "similarity": 0.9}]
db.usage_units = lambda kinds, delta: 1
db.storage_used_bytes = lambda: 1000
db.document_count = lambda: 1
db.get_profile = lambda: {"display_name": "Test", "email": "a@b.co"}
db.list_saved_answers = lambda: []
db.list_saved_diagrams = lambda: []
db.search_conversations = lambda t: CONVS
quota.check_ai_quota = lambda: None
storage.download = lambda p: BYTES[KIND]
llm.embed = lambda texts, query=False: [[0.1] * 768 for _ in texts]


def _gen(system, history, prompt):
    CALLS.append({"system": system, "prompt": prompt})
    return "Loop na oru code-a thirumba run panna use pannuvom [S1].\n```mermaid\nflowchart TD\n A[Start] --> B[End]\n```"


llm.generate = _gen
runpy.run_path(str(Path(__file__).resolve().parent.parent / "app.py"), run_name="__main__")
